import numpy as np
import pandas as pd

from ASFINT.Pipeline.workflow import _normalize_week, _write_csv
from ASFINT.Transform.Agenda_Processor import _docx_clean_org, _docx_decision
from ASFINT.Transform.FR_Processor import FR_ProcessorV2
from ASFINT.Transform.Reconciliation_Processor import Merge_FR_Agenda


def test_week_normalization():
    assert _normalize_week("FR 25_26 F02.xlsx") == "F2"
    assert _normalize_week("FR 25/26 S12") == "S12"


def test_agenda_org_cleanup():
    assert _docx_clean_org("Example Organization (Under Community Projects)*") == "Example Organization (Under Community Projects)"
    assert _docx_clean_org("Sample Club ($1,500)") == "Sample Club"


def test_agenda_decisions_keep_tabled_statuses_distinct():
    assert _docx_decision(["Motion to table for one week."])[0] == "Tabled"
    assert _docx_decision(["Motion to table indefinitely."])[0] == "Tabled Indefinitely"
    assert _docx_decision(["Motion to deny the request."])[0] == "Denied"
    assert _docx_decision(["Motion to partially approve for $250."])[0] == "Partially Approved"


def test_fr_repeated_header_is_treated_as_decision_table():
    raw = pd.DataFrame([
        ["Senate Finance Resolution 2025/2026 Fall Week 01/12/2026", None, None, None, None, None, None, None],
        ["Appx.", "Org Name", "Request Type", "Org Type (year)", "Amount Requested", "Funding Source", "Primary Contact", None],
        ["A", "Approved Org", "Contingency", "SAG 1", 100, "ASUC", "Person A", "a@example.edu"],
        ["B", "Tabled Org", "Contingency", "SAG 2", 200, "ASUC", "Person B", "b@example.edu"],
        [None, None, None, None, None, None, None, None],
        ["Appx.", "Org Name", "Request Type", "Org Type (year)", "Amount Requested", "Funding Source", "Primary Contact", None],
        ["A", "Approved Org", "Contingency", "SAG 1", 100, "ASUC", "Person A", "a@example.edu"],
        ["B", "Tabled Org", "Contingency", "SAG 2", "Tabled indefinitely", "ASUC", "Person B", "b@example.edu"],
    ])
    cleaned = next(iter(FR_ProcessorV2(raw, original_filename="FR F2.xlsx").values()))
    assert len(cleaned) == 2
    assert cleaned.loc[cleaned["Org Name"] == "Approved Org", "Committee Status"].iloc[0] == "Approved"
    assert cleaned.loc[cleaned["Org Name"] == "Approved Org", "Amount"].iloc[0] == 100
    assert cleaned.loc[cleaned["Org Name"] == "Tabled Org", "Committee Status"].iloc[0] == "Tabled Indefinitely"
    assert pd.isna(cleaned.loc[cleaned["Org Name"] == "Tabled Org", "Amount"].iloc[0])
    assert cleaned.loc[cleaned["Org Name"] == "Approved Org", "Email Address"].iloc[0] == "a@example.edu"


def test_merge_matches_on_request_type_not_org_name_only():
    fr = pd.DataFrame([
        {"Org Name": "Same Org", "Request Type": "Contingency", "Amount": 100, "Committee Status": "Approved"},
        {"Org Name": "Same Org", "Request Type": "Space Reservation", "Amount": 200, "Committee Status": "Approved"},
    ])
    agenda = pd.DataFrame([
        {"Org Name": "Same Org", "Request Type": "Contingency", "Amount": 90, "Committee Status": "Approved"},
        {"Org Name": "Same Org", "Request Type": "Space Reservation", "Amount": 180, "Committee Status": "Approved"},
    ])
    out, _ = Merge_FR_Agenda(fr, agenda, "2026-01-01")
    assert len(out) == 2
    assert set(out["Document Match Type"]) == {"Exact"}
    assert set(out["FR Amount"]) == {100, 200}
    assert set(out["Agenda Amount"]) == {90, 180}


def test_merge_preserves_disagreements():
    fr = pd.DataFrame([{"Org Name": "Example Org", "Request Type": "Contingency", "Amount": 700, "Committee Status": "Approved"}])
    agenda = pd.DataFrame([{"Org Name": "Example Org", "Request Type": "Contingency", "Amount": 100, "Committee Status": "Tabled"}])
    out, _ = Merge_FR_Agenda(fr, agenda, "2026-01-01")
    row = out.iloc[0]
    assert row["FR Amount"] == 700
    assert row["Agenda Amount"] == 100
    assert row["FR Committee Status"] == "Approved"
    assert row["Agenda Committee Status"] == "Tabled"


def test_merge_keeps_one_sided_rows():
    fr = pd.DataFrame([{"Org Name": "FR Org", "Request Type": "Contingency"}])
    agenda = pd.DataFrame([{"Org Name": "Agenda Org", "Request Type": "Contingency"}])
    out, _ = Merge_FR_Agenda(fr, agenda, "2026-01-01")
    assert set(out["Document Match Type"]) == {"FR Only", "Agenda Only"}


def test_merge_marks_fuzzy_names_without_rewriting_sources():
    fr = pd.DataFrame([{"Org Name": "Example Student Association", "Request Type": "Contingency"}])
    agenda = pd.DataFrame([{"Org Name": "Example Student Assoication", "Request Type": "Contingency"}])
    out, flags = Merge_FR_Agenda(fr, agenda, "2026-01-01")
    row = out.iloc[0]
    assert row["Document Match Type"] == "Fuzzy"
    assert row["FR Org Name"] == "Example Student Association"
    assert row["Agenda Org Name"] == "Example Student Assoication"
    assert row["Document Match Score"] >= 96
    assert "FUZZY_NAME_MATCH" in {flag["code"] for flag in flags}


def test_merge_matches_small_org_name_typo():
    fr = pd.DataFrame([{"Org Name": "Example Danco Club", "Request Type": "Space Reservation"}])
    agenda = pd.DataFrame([{"Org Name": "Example Dance Club", "Request Type": "Space Reservation"}])
    out, _ = Merge_FR_Agenda(fr, agenda, "2026-01-01")
    assert len(out) == 1
    assert out.iloc[0]["Document Match Type"] == "Fuzzy"


def test_agenda_parser_handles_senate_bill_and_public_comment(tmp_path):
    from docx import Document
    from ASFINT.Transform.Agenda_Processor import Agenda_Docx_Processor

    path = tmp_path / "agenda.docx"
    doc = Document()
    for line in [
        "Monday, January 12th, 2026",
        "FR 25/26 F1",
        "Sponsorship",
        "Example Club",
        "Motion to approve.",
        "Guest Announcements & Public Comment",
        "SR 25/26-001 Example Senate Bill",
        "Motion to approve as amended from Senator Example.",
        "Adjournment",
    ]:
        doc.add_paragraph(line)
    doc.save(path)

    out, _, _, _ = Agenda_Docx_Processor(path, "2026-01-12")
    assert set(out["Org Name"]) == {"Example Club", "SR 25/26-001 Example Senate Bill"}
    bill = out[out["Request Type"] == "Senate Bill"].iloc[0]
    assert bill["Committee Status"] == "Approved but dollar amount not listed"


def test_agenda_parser_extracts_senate_bill_added_by_motion(tmp_path):
    from docx import Document
    from ASFINT.Transform.Agenda_Processor import Agenda_Docx_Processor

    path = tmp_path / "agenda.docx"
    doc = Document()
    for line in [
        "Monday, January 12th, 2026",
        "Pending Business",
        "FR 25/26 S1",
        "Sponsorship",
        "Example Club",
        "Motion to approve.",
        "Motion to amend the agenda to add SR 25/26-002 Example Resolution from Senator Example.",
        "Motion to approve SR 25/26-002 from Senator Example.",
        "Senator Example asks the speaker about associated costs.",
        "Adjournment",
    ]:
        doc.add_paragraph(line)
    doc.save(path)

    out, _, _, _ = Agenda_Docx_Processor(path, "2026-01-12")
    assert "Senator Example asks the speaker about associated costs." not in set(out["Org Name"])
    bill = out[out["Request Type"] == "Senate Bill"].iloc[0]
    assert bill["Org Name"].startswith("SR 25/26-002")
    assert bill["Committee Status"] == "Approved but dollar amount not listed"


def test_senate_bills_match_by_resolution_number():
    fr = pd.DataFrame([{
        "Org Name": "SR 25/26-002",
        "Request Type": "Senate Bill",
    }])
    agenda = pd.DataFrame([{
        "Org Name": "SR 25/26-002 Example Resolution",
        "Request Type": "Senate Bill",
    }])
    out, flags = Merge_FR_Agenda(fr, agenda, "2026-01-12")
    assert len(out) == 1
    assert out.iloc[0]["Document Match Type"] == "Exact"
    assert "FUZZY_NAME_MATCH" not in {flag["code"] for flag in flags}


def test_senate_bill_scalar_decision_is_not_part_of_title(tmp_path):
    from docx import Document
    from ASFINT.Transform.Agenda_Processor import Agenda_Docx_Processor

    path = tmp_path / "agenda.docx"
    doc = Document()
    for line in [
        "Monday, January 12th, 2026",
        "Pending Business",
        "Senate Bills",
        "SR 25/26-003 Example Budget Resolution",
        "Approved as amended",
        "SR 25/26-004 Example Finance Resolution",
        "Approved",
        "Adjournment",
    ]:
        doc.add_paragraph(line)
    doc.save(path)

    out, _, _, _ = Agenda_Docx_Processor(path, "2026-01-12")
    assert list(out["Org Name"]) == [
        "SR 25/26-003 Example Budget Resolution",
        "SR 25/26-004 Example Finance Resolution",
    ]
    assert list(out["Committee Status"]) == ["Approved but dollar amount not listed", "Approved but dollar amount not listed"]


def test_agenda_numeric_after_denied_is_not_extra_outcome(tmp_path):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from ASFINT.Transform.Agenda_Processor import Agenda_Docx_Processor

    def add_numbered(text, level):
        paragraph = doc.add_paragraph(text)
        ppr = paragraph._p.get_or_add_pPr()
        numpr = OxmlElement("w:numPr")
        ilvl = OxmlElement("w:ilvl")
        ilvl.set(qn("w:val"), str(level))
        numid = OxmlElement("w:numId")
        numid.set(qn("w:val"), "1")
        numpr.append(ilvl)
        numpr.append(numid)
        ppr.append(numpr)

    path = tmp_path / "agenda.docx"
    doc = Document()
    doc.add_paragraph("Monday, January 12th, 2026")
    add_numbered("Pending Business", 0)
    add_numbered("FR 25/26 S1", 1)
    add_numbered("Space Reservation", 2)
    add_numbered("Example Organization A", 3)
    add_numbered("Denied", 4)
    add_numbered("1200", 4)
    add_numbered("Example Organization B", 3)
    add_numbered("2500", 4)
    add_numbered("Adjournment", 0)
    doc.save(path)

    out, _, _, flags = Agenda_Docx_Processor(path, "2026-01-12")

    denied = out[out["Org Name"] == "Example Organization A"].iloc[0]
    assert denied["Committee Status"] == "Denied"
    assert denied["Amount"] == 0
    assert "AGENDA_EXTRA_SCALAR" not in {flag["code"] for flag in flags}

    approved = out[out["Org Name"] == "Example Organization B"].iloc[0]
    assert approved["Committee Status"] == "Approved"
    assert float(approved["Amount"]) == 2500.0


def test_merge_recovers_combined_agenda_org_when_outcome_identifies_one_fr_row():
    fr = pd.DataFrame([
        {
            "Org Name": "Alpha Events",
            "Request Type": "Contingency",
            "Amount Requested": 125,
            "Amount": np.nan,
            "Committee Status": "Tabled Indefinitely",
        },
        {
            "Org Name": "Beta Student Society (BSS)",
            "Request Type": "Contingency",
            "Amount Requested": 650,
            "Amount": 420,
            "Committee Status": "Approved",
        },
    ])
    agenda = pd.DataFrame([{
        "Org Name": "Beta Student Society (BSS)Alpha Events",
        "Request Type": "Contingency",
        "Amount": 420,
        "Committee Status": "Approved",
    }])

    out, flags = Merge_FR_Agenda(fr, agenda, "2026-01-12")

    assert len(out) == 2
    matched = out[out["Agenda Org Name"].notna()].iloc[0]
    assert matched["Document Match Type"] == "Fuzzy"
    assert matched["FR Org Name"] == "Beta Student Society (BSS)"
    assert matched["Agenda Org Name"] == "Beta Student Society (BSS)Alpha Events"
    assert matched["Agenda Amount"] == 420

    unmatched = out[out["FR Org Name"] == "Alpha Events"].iloc[0]
    assert unmatched["Document Match Type"] == "FR Only"

    clerical = [flag for flag in flags if flag["code"] == "POTENTIAL_CLERICAL_ERROR"]
    assert len(clerical) == 1
    assert "Potential clerical error" in clerical[0]["message"]
    assert "recovered_match=Beta Student Society (BSS)" in clerical[0]["details"]


def test_merge_does_not_guess_when_combined_agenda_org_is_ambiguous():
    fr = pd.DataFrame([
        {
            "Org Name": "Alpha Club (AC)",
            "Request Type": "Contingency",
            "Amount": 100,
            "Committee Status": "Approved",
        },
        {
            "Org Name": "Beta Club",
            "Request Type": "Contingency",
            "Amount": 100,
            "Committee Status": "Approved",
        },
    ])
    agenda = pd.DataFrame([{
        "Org Name": "Alpha Club (AC)Beta Club",
        "Request Type": "Contingency",
        "Amount": 100,
        "Committee Status": "Approved",
    }])

    out, flags = Merge_FR_Agenda(fr, agenda, "2026-01-12")

    assert list(out["Document Match Type"]).count("FR Only") == 2
    assert list(out["Document Match Type"]).count("Agenda Only") == 1
    clerical = [flag for flag in flags if flag["code"] == "POTENTIAL_CLERICAL_ERROR"]
    assert len(clerical) == 1
    assert "no automatic match" in clerical[0]["details"]


def test_csv_output_formats_money_columns_to_two_decimals(tmp_path):
    frame = pd.DataFrame({
        "Amount Requested": [500, 110.5, np.nan, "N/A"],
        "Amount": [0, 12.345, 1000.0, np.nan],
        "Document Match Score": [100.0, 96.551724, np.nan, 93.3],
    })
    path = tmp_path / "output.csv"

    _write_csv(frame, path)

    lines = path.read_text().splitlines()
    assert lines[1] == "500.00,0.00,100.0"
    assert lines[2] == "110.50,12.35,96.551724"
    assert lines[3] == ",1000.00,"
    assert lines[4] == "N/A,,93.3"
