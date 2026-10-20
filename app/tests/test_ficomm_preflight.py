from __future__ import annotations

import asyncio
from datetime import date
from io import BytesIO
from pathlib import Path

import pandas as pd
import pytest
from docx import Document
from fastapi import HTTPException

from app.core.models import CollectionPeriod, Dataset, Ingestion
from app.services.ingestion_service import IngestionService


class FakeUploadFile:
    def __init__(self, filename: str, content: bytes, content_type: str):
        self.filename = filename
        self.content_type = content_type
        self._content = content

    async def read(self, size: int = -1) -> bytes:
        return self._content


def run(coro):
    return asyncio.run(coro)


def _ficomm_dataset(db) -> Dataset:
    dataset = Dataset(name="FICOMM", process_type="FICOMM")
    db.add(dataset)
    db.commit()
    db.refresh(dataset)
    return dataset


def _period(db, dataset_id: int, *, open_for_uploads: bool = True) -> CollectionPeriod:
    period = CollectionPeriod(
        dataset_id=dataset_id,
        label="Fall 2026",
        starts_on=date(2026, 8, 1),
        ends_on=date(2026, 12, 31),
        accept_uploads=open_for_uploads,
    )
    db.add(period)
    db.commit()
    db.refresh(period)
    return period


def _fr_bytes(meeting_date: str = "10/12/2026") -> bytes:
    raw = pd.DataFrame([
        [f"Senate Finance Resolution 2026/2027 {meeting_date}", None, None, None, None, None],
        ["Appx.", "Org Name", "Request Type", "Org Type (year)", "Amount Requested", "Funding Source"],
        ["A", "Example Organization", "Contingency Funding", "SAG 1", 100, "ASUC"],
    ])
    buffer = BytesIO()
    raw.to_excel(buffer, index=False, header=False)
    return buffer.getvalue()


def _agenda_bytes(week: str = "F7", meeting_date: str = "Monday, October 12th, 2026") -> bytes:
    document = Document()
    document.add_paragraph(meeting_date)
    document.add_paragraph(f"FR 26/27 {week}")
    document.add_paragraph("Pending Business")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def test_ficomm_fr_upload_is_classified_and_assigned_to_open_period(db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.storage.DATA_ROOT", tmp_path)
    dataset = _ficomm_dataset(db)
    period = _period(db, dataset.id)

    ingestion = run(
        IngestionService(db).create(
            dataset.id,
            FakeUploadFile(
                "FR 26_27 F7.xlsx",
                _fr_bytes(),
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            ),
        )
    )

    assert ingestion.period_id == period.id
    assert ingestion.source_role == "FR"
    assert ingestion.meeting_week == "F7"
    assert ingestion.meeting_date == date(2026, 10, 12)
    assert Path(ingestion.raw_path).exists()


def test_ficomm_agenda_upload_is_classified_and_assigned_to_open_period(db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.storage.DATA_ROOT", tmp_path)
    dataset = _ficomm_dataset(db)
    period = _period(db, dataset.id)

    ingestion = run(
        IngestionService(db).create(
            dataset.id,
            FakeUploadFile(
                "Agenda 10-12-2026.docx",
                _agenda_bytes(),
                "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            ),
        )
    )

    assert ingestion.period_id == period.id
    assert ingestion.source_role == "AGENDA"
    assert ingestion.meeting_week == "F7"
    assert ingestion.meeting_date == date(2026, 10, 12)


def test_ficomm_rejects_unsupported_file_before_creating_ingestion(db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.storage.DATA_ROOT", tmp_path)
    dataset = _ficomm_dataset(db)
    _period(db, dataset.id)

    with pytest.raises(HTTPException) as exc:
        run(
            IngestionService(db).create(
                dataset.id,
                FakeUploadFile("notes.pdf", b"not a ficomm source", "application/pdf"),
            )
        )

    assert exc.value.status_code == 422
    assert db.query(Ingestion).count() == 0
    assert not (tmp_path / "raw").exists()


def test_ficomm_rejects_upload_outside_open_period(db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.storage.DATA_ROOT", tmp_path)
    dataset = _ficomm_dataset(db)
    _period(db, dataset.id, open_for_uploads=False)

    with pytest.raises(HTTPException) as exc:
        run(
            IngestionService(db).create(
                dataset.id,
                FakeUploadFile(
                    "FR 26_27 F7.xlsx",
                    _fr_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        )

    assert exc.value.status_code == 422
    assert "No open collection period" in exc.value.detail
    assert db.query(Ingestion).count() == 0


def test_ficomm_rejects_week_that_conflicts_with_period(db, tmp_path, monkeypatch):
    monkeypatch.setattr("app.services.storage.DATA_ROOT", tmp_path)
    dataset = _ficomm_dataset(db)
    _period(db, dataset.id)

    with pytest.raises(HTTPException) as exc:
        run(
            IngestionService(db).create(
                dataset.id,
                FakeUploadFile(
                    "FR 26_27 S7.xlsx",
                    _fr_bytes(),
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            )
        )

    assert exc.value.status_code == 422
    assert "does not match collection period" in exc.value.detail
    assert db.query(Ingestion).count() == 0
